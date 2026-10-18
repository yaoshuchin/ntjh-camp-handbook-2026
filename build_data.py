import openpyxl, docx, json, re, os, sys

def mask_name(name):
    if not name or not isinstance(name, str):
        return ''
    name = name.strip()
    if name in ['—', '-', '無']:
        return '—'
    suffix = ''
    if '(素)' in name:
        name = name.replace('(素)', '').strip()
        suffix = '(素)'
    elif '（素）' in name:
        name = name.replace('（素）', '').strip()
        suffix = '(素)'
    if len(name) <= 1:
        return name + suffix
    elif len(name) == 2:
        return name[0] + 'O' + suffix
    else:
        return name[0] + 'O' + name[2:] + suffix

def mask_person_entry(text):
    if not text or text in ['—', '-', '無']:
        return '—'
    text = text.strip()
    m = re.match(r'^(\d+號)?\s*(.+)$', text)
    if m:
        num = m.group(1) or ''
        nm = m.group(2) or ''
        return f'{num} {mask_name(nm)}'.strip()
    return mask_name(text)

def mask_member_list(text):
    if not text or text in ['—', '-', '無']:
        return '—'
    items = re.split(r'[、,，]\s*', text.strip())
    masked = [mask_person_entry(it) for it in items if it.strip()]
    return '、'.join(masked)

# 1. Load Excel
excel_path = '南投國中115二年級隔宿手寫資料電子檔(完整試算表) (1).xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

# Mentors and school staff
mentors_raw = {
    '201': '黃麗月',
    '202': '黃惠君',
    '203': '林宜玟',
    '204': '李佩儒',
    '205': '林富銓',
    '206': '許招玉',
    '207': '黃瓊慧',
}
mentors_masked = {k: mask_name(v) for k, v in mentors_raw.items()}

# Parse Students
ws_students = wb['全校學生清冊']
students = []
for r in range(5, ws_students.max_row + 1):
    c_class = ws_students.cell(r, 1).value
    if not c_class:
        continue
    c_teacher = ws_students.cell(r, 2).value or ''
    c_seat = ws_students.cell(r, 3).value or 0
    c_name = ws_students.cell(r, 4).value or ''
    c_gender = ws_students.cell(r, 5).value or ''
    c_tent = ws_students.cell(r, 6).value or ''
    c_leader = ws_students.cell(r, 7).value or ''
    c_squad = ws_students.cell(r, 8).value or ''
    c_role = ws_students.cell(r, 9).value or ''
    c_diet = ws_students.cell(r, 10).value or '葷食'

    students.append({
        'id': len(students) + 1,
        'class': str(c_class).strip(),
        'teacher': mask_name(str(c_teacher).strip()),
        'seat': int(c_seat),
        'name': mask_name(str(c_name).strip()),
        'gender': str(c_gender).strip(),
        'tent': str(c_tent).strip(),
        'isLeader': '室長' in str(c_leader),
        'leaderLabel': '室長 (負責人)' if '室長' in str(c_leader) else '成員',
        'squad': str(c_squad).strip(),
        'role': str(c_role).strip(),
        'diet': str(c_diet).strip(),
        'isVeg': '素' in str(c_diet)
    })

# Parse Tents
ws_tents = wb['帳篷住宿分配總表']
tents = []
for r in range(5, ws_tents.max_row + 1):
    c_class = ws_tents.cell(r, 1).value
    if not c_class or '全體' in str(c_class) or '合計' in str(c_class):
        continue
    gender = str(ws_tents.cell(r, 2).value or '').strip()
    tent_id = str(ws_tents.cell(r, 3).value or '').strip()
    l_seat = ws_tents.cell(r, 4).value or 0
    l_name = str(ws_tents.cell(r, 5).value or '').strip()
    count_str = str(ws_tents.cell(r, 6).value or '').strip()
    members_raw = str(ws_tents.cell(r, 7).value or '').strip()
    notes = str(ws_tents.cell(r, 8).value or '').strip()

    tents.append({
        'id': len(tents) + 1,
        'class': str(c_class).strip(),
        'gender': gender,
        'tentName': tent_id,
        'leaderSeat': int(l_seat) if isinstance(l_seat, (int, float)) and l_seat > 0 else 0,
        'leaderName': mask_name(l_name),
        'count': count_str,
        'members': mask_member_list(members_raw),
        'notes': notes if notes not in ['—', '-'] else ''
    })

# Parse Squads
ws_squads = wb['野炊烤肉小隊編組表']
squads = []
for r in range(5, ws_squads.max_row + 1):
    c_class = ws_squads.cell(r, 1).value
    if not c_class or '全校' in str(c_class) or '合計' in str(c_class):
        continue
    sq_num = str(ws_squads.cell(r, 2).value or '').strip()
    sq_cnt = str(ws_squads.cell(r, 3).value or '').strip()
    leader = str(ws_squads.cell(r, 4).value or '').strip()
    deputy = str(ws_squads.cell(r, 5).value or '').strip()
    chef1 = str(ws_squads.cell(r, 6).value or '').strip()
    chef2 = str(ws_squads.cell(r, 7).value or '').strip()
    fire = str(ws_squads.cell(r, 8).value or '').strip()
    water = str(ws_squads.cell(r, 9).value or '').strip()
    affairs = str(ws_squads.cell(r, 10).value or '').strip()
    members = str(ws_squads.cell(r, 11).value or '').strip()

    squads.append({
        'id': len(squads) + 1,
        'class': str(c_class).strip(),
        'squadName': sq_num,
        'count': sq_cnt,
        'leader': mask_person_entry(leader),
        'deputy': mask_person_entry(deputy),
        'chef1': mask_person_entry(chef1),
        'chef2': mask_person_entry(chef2),
        'fire': mask_person_entry(fire),
        'water': mask_person_entry(water),
        'affairs': mask_person_entry(affairs),
        'members': mask_member_list(members),
    })

# Parse Word doc
doc = docx.Document('南投國中115學年度二年級(尖山埤)1(名間親子).docx')
day1_schedule = []
for row in doc.tables[0].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    day1_schedule.append({
        'time': cells[0],
        'activity': cells[1],
        'note': cells[2]
    })

day2_schedule = []
for row in doc.tables[1].rows[1:]:
    cells = [c.text.strip() for c in row.cells]
    day2_schedule.append({
        'time': cells[0],
        'activity': cells[1],
        'note': cells[2]
    })

# Save JSON representation to inspect
data_obj = {
    'stats': {
        'totalStudents': len(students),
        'maleStudents': len([s for s in students if s['gender'] == '男']),
        'femaleStudents': len([s for s in students if s['gender'] == '女']),
        'vegetarianStudents': len([s for s in students if s['isVeg']]),
        'totalTents': len(tents),
        'maleTents': len([t for t in tents if '男' in t['gender']]),
        'femaleTents': len([t for t in tents if '女' in t['gender']]),
        'totalSquads': len(squads),
    },
    'mentors': mentors_masked,
    'day1_schedule': day1_schedule,
    'day2_schedule': day2_schedule,
    'students_count': len(students),
    'tents_count': len(tents),
    'squads_count': len(squads)
}

print('Parsed successfully:')
print(json.dumps(data_obj, ensure_ascii=False, indent=2))
